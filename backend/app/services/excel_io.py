"""Excel import/export utilities for Companies, Persons, Positions, Shareholdings."""
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.security import hash_id_card
from app.models import Company, Person, Position, Shareholding


# Header definitions for both template generation and parsing.
COMPANY_HEADERS = [
    "name", "credit_code", "short_name", "legal_representative",
    "reg_capital", "established_date", "reg_address",
    "business_scope", "company_type", "parent_company_name", "remark",
]
PERSON_HEADERS = [
    "name", "id_card", "nationality", "phone", "email", "gender", "remark",
]
POSITION_HEADERS = [
    "person_name", "person_id_card", "company_name",
    "position_type", "start_date", "end_date", "document_no", "remark",
]
SHAREHOLDING_HEADERS = [
    "company_name", "shareholder_type", "shareholder_name", "shareholder_id_card_or_credit_code",
    "ratio", "amount", "currency", "contribute_method", "contribute_date", "remark",
]

TEMPLATES = {
    "companies": COMPANY_HEADERS,
    "persons": PERSON_HEADERS,
    "positions": POSITION_HEADERS,
    "shareholdings": SHAREHOLDING_HEADERS,
}


def build_template(kind: str) -> bytes:
    if kind not in TEMPLATES:
        raise ValueError(f"未知模板类型: {kind}")
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    ws.append(TEMPLATES[kind])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_date(value: Any):
    if value in (None, ""):
        return None
    if isinstance(value, (date, datetime)):
        return value if isinstance(value, date) and not isinstance(value, datetime) else value.date()
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_float(value: Any):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _rows(file_bytes: bytes) -> List[Dict[str, Any]]:
    wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else "" for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if all(v in (None, "") for v in raw):
            continue
        out.append({headers[i]: raw[i] for i in range(len(headers)) if headers[i]})
    return out


def import_companies(db: Session, file_bytes: bytes) -> Dict[str, int]:
    created = updated = 0
    rows = _rows(file_bytes)
    # First pass: create / update without parent links.
    name_to_id: Dict[str, int] = {}
    pending_parent: Dict[str, str] = {}
    for r in rows:
        name = (r.get("name") or "").strip()
        if not name:
            continue
        company = db.query(Company).filter(Company.name == name).first()
        if company is None:
            company = Company(name=name)
            db.add(company)
            created += 1
        else:
            updated += 1
        company.credit_code = (r.get("credit_code") or None) or company.credit_code
        company.short_name = r.get("short_name") or company.short_name or ""
        company.legal_representative = r.get("legal_representative") or company.legal_representative or ""
        company.reg_capital = _parse_float(r.get("reg_capital")) or company.reg_capital
        company.established_date = _parse_date(r.get("established_date")) or company.established_date
        company.reg_address = r.get("reg_address") or company.reg_address or ""
        company.business_scope = r.get("business_scope") or company.business_scope or ""
        company.company_type = r.get("company_type") or company.company_type or ""
        company.remark = r.get("remark") or company.remark or ""
        db.flush()
        name_to_id[name] = company.id
        parent_name = (r.get("parent_company_name") or "").strip()
        if parent_name:
            pending_parent[name] = parent_name
    # Resolve parent links.
    for child, parent in pending_parent.items():
        parent_id = name_to_id.get(parent) or (
            db.query(Company.id).filter(Company.name == parent).scalar()
        )
        if parent_id:
            db.query(Company).filter(Company.id == name_to_id[child]).update(
                {"parent_company_id": parent_id}
            )
    db.commit()
    return {"created": created, "updated": updated}


def import_persons(db: Session, file_bytes: bytes) -> Dict[str, int]:
    created = updated = 0
    for r in _rows(file_bytes):
        name = (r.get("name") or "").strip()
        id_card = str(r.get("id_card") or "").strip()
        if not name or not id_card:
            continue
        h = hash_id_card(id_card)
        person = db.query(Person).filter(Person.id_card_hash == h).first()
        if person is None:
            person = Person(
                name=name,
                id_card_hash=h,
                id_card_last4=id_card[-4:],
            )
            db.add(person)
            created += 1
        else:
            updated += 1
        person.name = name
        person.nationality = r.get("nationality") or person.nationality or "中国"
        person.phone = r.get("phone") or person.phone or ""
        person.email = r.get("email") or person.email or ""
        person.gender = r.get("gender") or person.gender or ""
        person.remark = r.get("remark") or person.remark or ""
    db.commit()
    return {"created": created, "updated": updated}


def import_positions(db: Session, file_bytes: bytes) -> Dict[str, int]:
    created = skipped = 0
    for r in _rows(file_bytes):
        person_name = (r.get("person_name") or "").strip()
        id_card = str(r.get("person_id_card") or "").strip()
        company_name = (r.get("company_name") or "").strip()
        position_type = (r.get("position_type") or "").strip()
        if not (person_name and id_card and company_name and position_type):
            skipped += 1
            continue
        h = hash_id_card(id_card)
        person = db.query(Person).filter(Person.id_card_hash == h).first()
        company = db.query(Company).filter(Company.name == company_name).first()
        if not person or not company:
            skipped += 1
            continue
        pos = Position(
            person_id=person.id,
            company_id=company.id,
            position_type=position_type,
            start_date=_parse_date(r.get("start_date")),
            end_date=_parse_date(r.get("end_date")),
            document_no=r.get("document_no") or "",
            remark=r.get("remark") or "",
        )
        db.add(pos)
        created += 1
    db.commit()
    return {"created": created, "skipped": skipped}


def import_shareholdings(db: Session, file_bytes: bytes) -> Dict[str, int]:
    created = skipped = 0
    for r in _rows(file_bytes):
        company_name = (r.get("company_name") or "").strip()
        s_type = (r.get("shareholder_type") or "").strip().lower()
        s_name = (r.get("shareholder_name") or "").strip()
        s_key = str(r.get("shareholder_id_card_or_credit_code") or "").strip()
        if not (company_name and s_type and s_name):
            skipped += 1
            continue
        company = db.query(Company).filter(Company.name == company_name).first()
        if not company:
            skipped += 1
            continue
        if s_type == "person":
            if not s_key:
                skipped += 1
                continue
            person = db.query(Person).filter(Person.id_card_hash == hash_id_card(s_key)).first()
            if not person:
                skipped += 1
                continue
            shareholder_id = person.id
        elif s_type == "company":
            other = db.query(Company).filter(Company.name == s_name).first()
            if not other:
                skipped += 1
                continue
            shareholder_id = other.id
        else:
            skipped += 1
            continue
        sh = Shareholding(
            company_id=company.id,
            shareholder_type=s_type,
            shareholder_id=shareholder_id,
            ratio=_parse_float(r.get("ratio")),
            amount=_parse_float(r.get("amount")),
            currency=r.get("currency") or "CNY",
            contribute_method=r.get("contribute_method") or "货币",
            contribute_date=_parse_date(r.get("contribute_date")),
            remark=r.get("remark") or "",
        )
        db.add(sh)
        created += 1
    db.commit()
    return {"created": created, "skipped": skipped}
